import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 51, 102)

def create_server_documentation():
    doc = docx.Document()

    # Title
    title = doc.add_heading('Documentación del Servidor Intranet (Ubuntu 24.04)', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph('Este documento detalla la configuración y arquitectura del servidor de producción y QA para el proyecto Intranet.')

    # 1. Componentes Instalados
    add_heading(doc, '1. Componentes Instalados', 1)
    doc.add_paragraph('El servidor utiliza las siguientes tecnologías principales:')
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run('Sistema Operativo:').bold = True
    ul.add_run(' Ubuntu 24.04 LTS')
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run('Base de Datos:').bold = True
    ul.add_run(' PostgreSQL')
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run('Lenguaje Base:').bold = True
    ul.add_run(' Python 3 (entornos virtuales venv)')
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run('Servidor Web / Proxy Inverso:').bold = True
    ul.add_run(' Apache2')
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run('Framework Web:').bold = True
    ul.add_run(' FastAPI (ejecutado con Uvicorn)')
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run('ORM:').bold = True
    ul.add_run(' SQLAlchemy')
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run('Otras librerías:').bold = True
    ul.add_run(' psycopg2-binary, passlib, bcrypt, jinja2, pandas, openpyxl, python-jose, python-multipart')

    # 2. Rutas y Carpetas
    add_heading(doc, '2. Rutas y Carpetas Principales', 1)
    
    add_heading(doc, 'Entorno de Producción (PROD)', 2)
    p = doc.add_paragraph()
    p.add_run('Ruta Base: ').bold = True
    p.add_run('/var/www/intranet_prod\n')
    p.add_run('Archivos Estáticos: ').bold = True
    p.add_run('/var/www/intranet_prod/static\n')
    p.add_run('Plantillas HTML: ').bold = True
    p.add_run('/var/www/intranet_prod/templates\n')
    p.add_run('Entorno Virtual: ').bold = True
    p.add_run('/var/www/intranet_prod/venv\n')
    
    add_heading(doc, 'Entorno de Pruebas (QA)', 2)
    p = doc.add_paragraph()
    p.add_run('Ruta Base: ').bold = True
    p.add_run('/var/www/intranet_qa\n')
    p.add_run('Entorno Virtual: ').bold = True
    p.add_run('/var/www/intranet_qa/venv\n')

    # 3. Configuración de Base de Datos
    add_heading(doc, '3. Configuración de Bases de Datos', 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Parámetro'
    hdr_cells[1].text = 'Valor'
    
    row_cells = table.add_row().cells
    row_cells[0].text = 'Motor'
    row_cells[1].text = 'PostgreSQL'

    row_cells = table.add_row().cells
    row_cells[0].text = 'Usuario DB'
    row_cells[1].text = 'intranet_user'

    row_cells = table.add_row().cells
    row_cells[0].text = 'Contraseña DB'
    row_cells[1].text = 'Redes2010'
    
    row_cells = table.add_row().cells
    row_cells[0].text = 'Base de Datos PROD'
    row_cells[1].text = 'intranet_prod'
    
    row_cells = table.add_row().cells
    row_cells[0].text = 'Base de Datos QA'
    row_cells[1].text = 'intranet_qa'

    # 4. Servicios y Puertos
    add_heading(doc, '4. Servicios y Puertos', 1)
    table2 = doc.add_table(rows=1, cols=3)
    table2.style = 'Table Grid'
    hdr_cells2 = table2.rows[0].cells
    hdr_cells2[0].text = 'Servicio'
    hdr_cells2[1].text = 'Puerto Interno'
    hdr_cells2[2].text = 'Puerto Expuesto (Apache)'
    
    row_cells2 = table2.add_row().cells
    row_cells2[0].text = 'FastAPI - Producción (uvicorn)'
    row_cells2[1].text = '8000'
    row_cells2[2].text = '80 (pow-intranet)'

    row_cells2 = table2.add_row().cells
    row_cells2[0].text = 'FastAPI - QA (uvicorn)'
    row_cells2[1].text = '8001'
    row_cells2[2].text = '8080 (pow-intranet-qa)'

    # 5. Claves y Variables de Entorno (URLs)
    add_heading(doc, '5. Claves y Archivos de Configuración', 1)
    doc.add_paragraph('Cada entorno requiere un archivo .env en la raíz del proyecto. Estos contienen la configuración de la cadena de conexión.')
    
    p = doc.add_paragraph()
    p.add_run('Archivo en PROD (.env):').bold = True
    doc.add_paragraph('DATABASE_URL=postgresql://intranet_user:Redes2010@localhost:5432/intranet_prod', style='Quote')
    
    p = doc.add_paragraph()
    p.add_run('Archivo en QA (.env):').bold = True
    doc.add_paragraph('DATABASE_URL=postgresql://intranet_user:Redes2010@localhost:5432/intranet_qa', style='Quote')

    # 6. Servicios Systemd
    add_heading(doc, '6. Servicios Systemd', 1)
    doc.add_paragraph('Los servicios están configurados para auto-inicio y manejo de procesos en background:')
    ul2 = doc.add_paragraph(style='List Bullet')
    ul2.add_run('Servicio PROD:').bold = True
    ul2.add_run(' /etc/systemd/system/intranet_prod.service')
    ul2 = doc.add_paragraph(style='List Bullet')
    ul2.add_run('Servicio QA:').bold = True
    ul2.add_run(' /etc/systemd/system/intranet_qa.service')
    
    doc.add_paragraph('Comandos útiles para gestión de los servicios:')
    doc.add_paragraph('systemctl start|stop|restart intranet_prod', style='Quote')
    doc.add_paragraph('systemctl start|stop|restart intranet_qa', style='Quote')

    # Guardar
    doc.save('Documentacion_Servidor.docx')
    print('Documento guardado con éxito como Documentacion_Servidor.docx')

if __name__ == "__main__":
    create_server_documentation()
